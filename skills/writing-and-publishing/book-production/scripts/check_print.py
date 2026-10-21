#!/usr/bin/env python3
"""
check_print.py - verify a print interior PDF or a Word manuscript before handoff.

    python check_print.py INTERIOR.pdf --trim 6x9 [--bleed] [--png-dir DIR] [--expect "Chapter title" ...]
    python check_print.py MANUSCRIPT.docx

PDF checks (KDP / IngramSpark interior rules):
  page size = trim (+0.125in bleed on outer edges if --bleed), page count 24-828,
  every font embedded, chapter openers on odd pages, text inside the live area
  (KDP gutter for this page count, 0.25in outside minimum), image resolution >= 300 ppi,
  leftover placeholders ("TODO", "ISBN: to be assigned", "000" TOC numbers, lorem),
  blank pages, and it renders proof PNGs (title, contents, first opener, a mid-book
  spread, last page) so the pages get looked at, not only measured.

DOCX checks (Scribe lessons): heading outline, comment markers, tracked changes,
  blank-paragraph ratio, placeholders, fused words, emoji, mixed quote styles,
  page size and footer page fields.
Exit code 1 when a must-fix problem is found.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import bookcore as bc  # noqa: E402
from build_book import TRIMS, kdp_min_gutter  # noqa: E402

PLACE = re.compile(r"\b(TODO|TBD|FIXME|lorem ipsum|to be assigned|NEW SECTION|XX+)\b|\?\?\?", re.I)


def check_pdf(path: Path, trim_key: str, bleed: bool, png_dir: Path | None, expect: list[str], platform: str = "kdp"):
    import fitz
    must, should, info = [], [], []
    doc = fitz.open(str(path))
    n = len(doc)
    w, h = TRIMS[trim_key]
    exp_w = (w + (0.125 if bleed else 0)) * 72
    exp_h = (h + (0.25 if bleed else 0)) * 72
    sizes = {(round(p.rect.width, 1), round(p.rect.height, 1)) for p in doc}
    if any(abs(sw - exp_w) > 1 or abs(sh - exp_h) > 1 for sw, sh in sizes):
        must.append(f"page size {sizes} pt does not match trim {trim_key}{' + bleed' if bleed else ''} "
                    f"({exp_w:.1f} x {exp_h:.1f} pt)")
    info.append(f"pages: {n}  |  size: {w} x {h} in{' + bleed' if bleed else ''}")
    if n < 24:
        (must if platform in ("kdp", "ingram") else should).append(
            f"{n} pages: under KDP's 24-page minimum" + ("" if platform in ("kdp", "ingram") else
            "; a print shop can still bind it (saddle stitch suits short books; page count must be divisible by 4)"))
    if n > 828:
        must.append(f"{n} pages: over the KDP paperback maximum (828)")
    if n % 2:
        should.append("odd page count; printers add a blank at the end (fine, but intentional is better)")
    # fonts
    unembedded = set()
    names = set()
    for p in doc:
        for f in p.get_fonts(full=True):
            names.add(f[3].split("+")[-1])
            if f[1] in ("n/a", "") or not f[0]:
                unembedded.add(f[3])
    if unembedded:
        must.append(f"fonts not embedded: {', '.join(sorted(unembedded))}")
    info.append(f"fonts: {', '.join(sorted(names))}")
    # openers
    toc = doc.get_toc(simple=True)
    even_openers = [f"{t} (p{pg})" for lvl, t, pg in toc if lvl == 1 and pg % 2 == 0 and not t.lower().startswith(("copyright",))]
    if even_openers:
        should.append("sections opening on a verso (even page): " + "; ".join(even_openers[:8]))
    info.append(f"bookmarks: {len(toc)}")
    # live area
    gutter = kdp_min_gutter(n) * 72
    outside = 0.25 * 72
    bleed_off = 0.125 * 72 if bleed else 0
    out_of_area, blanks, low_res = [], [], []
    full_text = []
    for i, p in enumerate(doc, 1):
        txt = p.get_text()
        full_text.append(txt)
        if not txt.strip() and not p.get_images():
            blanks.append(i)
            continue
        recto = i % 2 == 1
        left_min = (gutter if recto else outside)
        right_max = p.rect.width - bleed_off - (outside if recto else gutter)
        for b in p.get_text("blocks"):
            x0, y0, x1, y1, t = b[0], b[1], b[2], b[3], b[4]
            if not t.strip():
                continue
            if x0 < left_min - 1 or x1 > right_max + 1:
                out_of_area.append(f"p{i}: '{t.strip()[:30]}' x {x0:.0f}-{x1:.0f}pt (allowed {left_min:.0f}-{right_max:.0f})")
                break
        for img in p.get_image_info(xrefs=True):
            bw = img["bbox"][2] - img["bbox"][0]
            if bw <= 0:
                continue
            ppi = img["width"] / (bw / 72)
            if ppi < 300:
                low_res.append(f"p{i}: {ppi:.0f} ppi")
    if out_of_area:
        must.append(f"text outside the safe area on {len(out_of_area)} page(s): " + "; ".join(out_of_area[:5]))
    if low_res:
        should.append("images under 300 ppi: " + ", ".join(low_res[:8]))
    info.append(f"blank pages: {len(blanks)} {blanks[:20]}")
    joined = "\n".join(full_text)
    ph = sorted({m.group(0) for m in PLACE.finditer(joined)})
    if ph:
        must.append("placeholders in the text: " + ", ".join(ph))
    if "ISBN" not in joined:
        should.append("no ISBN on the copyright page: fine for private copies, required to sell through stores or KDP/IngramSpark")
    if re.search(r"(?m)^\s*000\s*$", joined):
        must.append("contents page still shows 000 placeholder page numbers")
    for e in expect:
        if e not in joined:
            must.append(f"expected text missing from PDF: {e!r}")
    fused = sorted(set(bc.combined_word_suspects(joined)))
    if fused:
        should.append("possible fused words: " + ", ".join(fused[:12]))
    if png_dir:
        png_dir.mkdir(parents=True, exist_ok=True)
        picks = {"01-title": 1}
        toc_page = next((pg for lvl, t, pg in toc if t.lower() in ("contents", "table of contents")), None)
        if toc_page:
            picks["02-contents"] = toc_page
        body = [pg for lvl, t, pg in toc if lvl == 1 and re.match(r"chapter|part|introduction", t, re.I)]
        if body:
            picks["03-first-opener"] = body[0]
        mid = n // 2 if (n // 2) % 2 == 0 else n // 2 + 1
        picks["04-spread-mid"] = (mid, min(mid + 1, n))
        picks["05-last"] = max((i for i in range(1, n + 1) if i not in blanks), default=n)
        for name, pg in picks.items():
            pages = pg if isinstance(pg, tuple) else (pg,)
            pixs = [doc[x - 1].get_pixmap(dpi=90) for x in pages]
            if len(pixs) == 1:
                pixs[0].save(str(png_dir / f"{name}-p{pages[0]}.png"))
            else:
                from PIL import Image
                ims = [Image.frombytes("RGB", (px.width, px.height), px.samples) for px in pixs]
                gap = 12
                sheet = Image.new("RGB", (sum(i.width for i in ims) + gap, max(i.height for i in ims)), (120, 120, 120))
                sheet.paste(ims[0], (0, 0))
                sheet.paste(ims[1], (ims[0].width + gap, 0))
                sheet.save(png_dir / f"{name}-p{pages[0]}-{pages[1]}.png")
        info.append(f"proof images: {png_dir}")
    return must, should, info


def check_docx(path: Path):
    from docx import Document
    must, should, info = [], [], []
    paras, facts = bc.read_docx(path)
    doc = Document(str(path))
    sec = doc.sections[0]
    info.append(f"page: {sec.page_width.inches:.2f} x {sec.page_height.inches:.2f} in, margins L{sec.left_margin.inches:.2f} "
                f"R{sec.right_margin.inches:.2f} T{sec.top_margin.inches:.2f} B{sec.bottom_margin.inches:.2f}")
    if facts["comments"]:
        must.append(f"{len(facts['comments'])} reviewer comment(s) still in the file")
    if facts["insertions"] or facts["deletions"]:
        must.append(f"tracked changes still in the file ({facts['insertions']} ins, {facts['deletions']} del)")
    blank = sum(1 for p in paras if not p.text.strip())
    ratio = blank / max(len(paras), 1)
    (should if ratio > 0.15 else info).append(f"blank paragraphs: {blank}/{len(paras)} ({ratio:.0%})"
                                              + (" - spot-check for leftover split artifacts" if ratio > 0.15 else ""))
    text = "\n".join(p.text for p in paras)
    ph = sorted({m.group(0) for m in PLACE.finditer(text)})
    if ph:
        must.append("placeholders in the text: " + ", ".join(ph))
    fused = sorted(set(bc.combined_word_suspects(text)))
    if fused:
        should.append("possible fused words: " + ", ".join(fused[:12]))
    u = bc.unicode_report(text)
    if u["emoji"]:
        should.append(f"{u['emoji']} emoji characters (most print fonts have no glyphs)")
    if u["straight_double"] and u["curly_double"]:
        should.append(f"mixed double quotes: {u['straight_double']} straight, {u['curly_double']} curly")
    info.append(f"em dashes in text: {u['em_dash']} (author's call)")
    footers = "".join(p._p.xml for s in doc.sections for p in s.footer.paragraphs)
    info.append("footer PAGE field: " + ("yes" if "PAGE" in footers else "no (fine if the author adds numbers in Word)"))
    outline = [("  " * (p.level - 1)) + p.text.strip()[:70] for p in paras if p.level and p.text.strip()]
    info.append(f"outline ({len(outline)} headings):\n    " + "\n    ".join(outline[:60]))
    return must, should, info


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--trim", default="6x9")
    ap.add_argument("--bleed", action="store_true")
    ap.add_argument("--png-dir")
    ap.add_argument("--expect", nargs="*", default=[])
    ap.add_argument("--platform", default="kdp", choices=["kdp", "ingram", "printshop"])
    a = ap.parse_args()
    f = Path(a.file)
    if f.suffix.lower() == ".pdf":
        png = Path(a.png_dir) if a.png_dir else bc.work_root("proofs", f.stem)
        must, should, info = check_pdf(f, a.trim.lower().replace(" ", ""), a.bleed, png, a.expect, a.platform)
    else:
        must, should, info = check_docx(f)
    print(f"# Print check: {f.name}\n")
    for i in info:
        print(f"- {i}")
    print(f"\n## Must fix ({len(must)})")
    for m in must:
        print(f"- {m}")
    print(f"\n## Should fix ({len(should)})")
    for s in should:
        print(f"- {s}")
    sys.exit(1 if must else 0)


if __name__ == "__main__":
    main()
